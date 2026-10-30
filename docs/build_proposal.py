"""Build an editable Word proposal from the maintained Markdown source.

Run: python docs/build_proposal.py
Requires python-docx. No model, network, or experimental data is used.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent


def inline(paragraph, text):
    for token in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(token)


def field(paragraph, instruction, placeholder=None):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    command = OxmlElement("w:instrText")
    command.set(qn("xml:space"), "preserve")
    command.text = instruction
    run._r.extend([begin, command])
    if placeholder:
        separator = OxmlElement("w:fldChar")
        separator.set(qn("w:fldCharType"), "separate")
        run._r.append(separator)
        text = OxmlElement("w:t")
        text.text = placeholder
        run._r.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Times New Roman"
        doc.styles[name].font.color.rgb = RGBColor.from_string("16324F")
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Heading 1"].font.size = Pt(18)
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.styles["Heading 3"].font.size = Pt(12)
    doc.core_properties.title = "Evidence-Grounded AI Assistance for Database Security Test Reporting and OWASP-Aligned Remediation"
    doc.core_properties.subject = "FPT research thesis proposal — user-specified RQ1, RQ2 and RQ3"
    doc.core_properties.author = "[Group members to be completed]"
    doc.core_properties.keywords = "database security, OWASP, evidence grounding, RAG, remediation"
    header = section.header.paragraphs[0]
    header.text = "FPT UNIVERSITY  |  RESEARCH THESIS PROPOSAL"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("607080")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Proposal — no experimental results claimed  |  ").font.size = Pt(9)
    field(footer, " PAGE ")

    for text in ("MINISTRY OF EDUCATION AND TRAINING", "FPT UNIVERSITY"):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
    doc.add_paragraph()
    p = doc.add_paragraph("CAPSTONE PROJECT DOCUMENT", "Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(doc.core_properties.title, "Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Research Proposal and Evaluation Plan")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in [
        ("Group members / IDs", "[To be completed]"),
        ("Supervisor", "[To be completed]"),
        ("Project code", "[To be completed]"),
        ("Submission date / location", "[To be completed]"),
        ("Status", "Proposal; experiments have not yet been conducted"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    doc.add_page_break()
    p = doc.add_paragraph("Table of contents", "Heading 2")
    field(doc.add_paragraph(), ' TOC \\o "1-3" \\h \\z \\u ', "Update this field in Word to generate the table of contents.")
    doc.add_paragraph("In Word: select all and update fields before submission. Title metadata remains to be completed. Figures and tables should be captioned and indexed when the final thesis is assembled.")
    doc.add_page_break()

    lines = (ROOT / "research-proposal.md").read_text().splitlines()
    # The cover already carries the first metadata block. Preserve all content
    # starting with the document-use statement.
    start = next(i for i, line in enumerate(lines) if line.startswith("## How this proposal"))
    index = start
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in row):
                    rows.append(row)
                index += 1
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Table Grid"
                for row_number, row in enumerate(rows):
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        inline(cell.paragraphs[0], text)
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(10)
                            if row_number == 0:
                                run.bold = True
                    if row_number == 0:
                        repeat = OxmlElement("w:tblHeader")
                        table.rows[0]._tr.get_or_add_trPr().append(repeat)
                doc.add_paragraph()
            continue
        match = re.match(r"^(#{1,3})\s+(.*)", line)
        if match:
            level = len(match.group(1))
            doc.add_paragraph(match.group(2), f"Heading {level}")
        elif re.match(r"^\d+\.\s", line):
            inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            inline(doc.add_paragraph(), line)
        index += 1
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    destination = ROOT / "research-proposal.docx"
    doc.save(destination)
    return destination


if __name__ == "__main__":
    print(build())
